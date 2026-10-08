"""Estilos de los libros de Excel. Puerto de `src/lib/export/styles.ts`.

La versión desplegada generaba los archivos en el navegador con
`write-excel-file`; aquí se generan en el servidor con `openpyxl`. La paleta,
las medidas, los formatos numéricos y la estructura de las hojas son los
mismos, de modo que los archivos salen idénticos a los que la empresa ya
conoce.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date as _date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta de la marca (sin «#», como la pide openpyxl).
XL = {
    "ink": "0A0A0A",
    "inkSoft": "3F3F46",
    "accent": "7A00CC",
    "accentSoft": "F3E9FF",
    "pending": "EB6834",
    "pendingSoft": "FDEFE8",
    "border": "D4D4D8",
    "borderSoft": "E4E4E7",
    "soft": "F4F4F5",
    "muted": "71717A",
    "white": "FFFFFF",
    "weekend": "EDEDED",
    "holiday": "E8F5EA",
    "positive": "166534",
    "negative": "B91C1C",
}

FORMATO_HORAS = "0.00"
FORMATO_DINERO = "#,##0.00"
FORMATO_FECHA = "dd.mm.yyyy"

_HAIR = Side(style="hair", color=XL["borderSoft"])
_THIN_BORDE = Side(style="thin", color=XL["border"])
_THIN_TINTA = Side(style="thin", color=XL["ink"])


@dataclass
class Celda:
    """Una celda con su valor y su estilo, tal como en el reporte original."""

    valor: object = None
    negrita: bool = False
    tamano: int = 10
    color: str | None = None
    fondo: str | None = None
    alineacion: str | None = None
    vertical: str | None = None
    ajustar: bool = False
    columnas: int = 1
    formato: str | None = None
    borde_inferior: str | None = None
    borde_superior: str | None = None
    alto: float | None = None


def relleno(cantidad: int) -> list:
    """Celdas vacías que completan una fila combinada."""
    return [None] * cantidad


def fila_vacia(ancho: int) -> list:
    return relleno(ancho)


def fila_titulo(valor: str, ancho: int) -> list:
    return [
        Celda(valor, negrita=True, tamano=16, color=XL["white"], fondo=XL["ink"],
              vertical="center", alto=34, columnas=ancho)
    ] + relleno(ancho - 1)


def fila_subtitulo(valor: str, ancho: int) -> list:
    return [
        Celda(valor, negrita=True, tamano=11, color=XL["white"], fondo=XL["accent"],
              vertical="center", alto=22, columnas=ancho)
    ] + relleno(ancho - 1)


def fila_leyenda(valor: str, ancho: int) -> list:
    return [
        Celda(valor, tamano=9, color=XL["muted"], vertical="center", alto=18,
              columnas=ancho)
    ] + relleno(ancho - 1)


def fila_seccion(valor: str, ancho: int) -> list:
    return [
        Celda(valor, negrita=True, tamano=11, color=XL["ink"], fondo=XL["soft"],
              borde_inferior="thin", vertical="center", alto=24, columnas=ancho)
    ] + relleno(ancho - 1)


def fila_encabezado(etiquetas: list) -> list:
    return [
        Celda(etiqueta, negrita=True, tamano=10, color=XL["white"], fondo=XL["ink"],
              alineacion="left", vertical="center", ajustar=True, alto=26)
        for etiqueta in etiquetas
    ]


def texto(valor, **opciones) -> Celda:
    opciones.setdefault("borde_inferior", "hair")
    return Celda(valor if valor is not None else "", **opciones)


def numero(valor, formato: str = FORMATO_HORAS, **opciones) -> Celda:
    opciones.setdefault("alineacion", "right")
    opciones.setdefault("borde_inferior", "hair")
    return Celda(valor, formato=formato, **opciones)


def dinero(valor, **opciones) -> Celda:
    return numero(valor, formato=FORMATO_DINERO, **opciones)


def fecha(iso: str, **opciones) -> Celda:
    opciones.setdefault("alineacion", "left")
    opciones.setdefault("borde_inferior", "hair")
    return Celda(_a_fecha(iso), formato=FORMATO_FECHA, **opciones)


def _a_fecha(iso: str):
    try:
        anio, mes, dia = (int(parte) for parte in str(iso)[:10].split("-"))
        return _date(anio, mes, dia)
    except (ValueError, TypeError):
        return iso


def etiqueta(valor: str, **opciones) -> Celda:
    """Etiqueta de un bloque clave-valor: vacaciones, indicadores."""
    opciones.setdefault("negrita", True)
    opciones.setdefault("color", XL["inkSoft"])
    return texto(valor, **opciones)


# Fila de totales: negrita, fondo suave y línea oscura arriba.
ESTILO_TOTAL = {
    "negrita": True,
    "fondo": XL["soft"],
    "borde_superior": "thin-tinta",
}


def total(valor=None, **opciones) -> Celda:
    combinado = dict(ESTILO_TOTAL)
    combinado.update(opciones)
    return texto(valor if valor is not None else "", **combinado)


def total_numero(valor, **opciones) -> Celda:
    combinado = dict(ESTILO_TOTAL)
    combinado.update(opciones)
    return numero(valor, **combinado)


def color_diferencia(valor: float) -> str | None:
    """Color del texto según el signo de una diferencia de horas."""
    if valor > 0:
        return XL["positive"]
    if valor < 0:
        return XL["negative"]
    return None


def nombre_hoja(valor: str) -> str:
    """Nombre de hoja válido para Excel: máximo 31 caracteres, sin símbolos."""
    return re.sub(r"[:\\/?*\[\]]", " ", valor or "")[:31] or "Hoja"


@dataclass
class Hoja:
    """Una hoja del libro: nombre, filas, anchos y filas congeladas."""

    nombre: str
    filas: list
    columnas: list = field(default_factory=list)
    filas_fijas: int = 0


def volcar(libro, hoja: Hoja) -> None:
    """Escribe una hoja completa en el libro de openpyxl."""
    destino = libro.create_sheet(title=nombre_hoja(hoja.nombre))

    for indice, ancho in enumerate(hoja.columnas, start=1):
        destino.column_dimensions[get_column_letter(indice)].width = ancho.get("width", 12)

    for numero_fila, fila in enumerate(hoja.filas, start=1):
        alto_fila = None
        # Las celdas `None` son las que cubre una combinada: se saltan.
        for numero_columna, celda in enumerate(fila, start=1):
            if celda is None:
                continue

            destino_celda = destino.cell(row=numero_fila, column=numero_columna)
            destino_celda.value = celda.valor

            destino_celda.font = Font(
                name="Calibri",
                bold=celda.negrita,
                size=celda.tamano,
                color=celda.color or "000000",
            )
            if celda.fondo:
                destino_celda.fill = PatternFill(
                    "solid", start_color=celda.fondo, end_color=celda.fondo
                )
            destino_celda.alignment = Alignment(
                horizontal=celda.alineacion,
                vertical=celda.vertical or "center",
                wrap_text=celda.ajustar,
            )
            if celda.formato:
                destino_celda.number_format = celda.formato

            abajo = (
                _HAIR
                if celda.borde_inferior == "hair"
                else _THIN_BORDE
                if celda.borde_inferior == "thin"
                else None
            )
            arriba = (
                _THIN_TINTA
                if celda.borde_superior == "thin-tinta"
                else _THIN_BORDE
                if celda.borde_superior == "thin"
                else None
            )
            if abajo or arriba:
                destino_celda.border = Border(bottom=abajo, top=arriba)

            if celda.columnas > 1:
                destino.merge_cells(
                    start_row=numero_fila,
                    start_column=numero_columna,
                    end_row=numero_fila,
                    end_column=numero_columna + celda.columnas - 1,
                )
            if celda.alto:
                alto_fila = max(alto_fila or 0, celda.alto)

        if alto_fila:
            destino.row_dimensions[numero_fila].height = alto_fila

    if hoja.filas_fijas:
        destino.freeze_panes = f"A{hoja.filas_fijas + 1}"
